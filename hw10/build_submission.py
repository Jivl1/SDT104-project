from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
REPORT_PATH = ROOT / "SDT104-hw10-illia.docx"
SCREENSHOT_DIR = ROOT / "screenshots"
REPO_URL = "https://github.com/jivl1/SDT104-project/tree/main/hw10"
LIVE_URL = "https://jivl1.github.io/SDT104-project/hw10/"

TASKS = [
    {
        "id": 1,
        "title": "Variables and Data Types",
        "file": ROOT / "task1.js",
        "learned": (
            "I practiced the difference between let and const, and used template literals "
            "to build readable strings with embedded values."
        ),
    },
    {
        "id": 2,
        "title": "Arrays and Loops",
        "file": ROOT / "task2.js",
        "learned": (
            "I worked with arrays, looped through items with forEach, transformed text with "
            "toUpperCase, and added a new value with push."
        ),
    },
    {
        "id": 3,
        "title": "Object Basics",
        "file": ROOT / "task3.js",
        "learned": (
            "I created an object with mixed property types and added a method that uses this "
            "to return a short summary of the hobby."
        ),
    },
    {
        "id": 4,
        "title": "Control Structures",
        "file": ROOT / "task4.js",
        "learned": (
            "I used an if / else if / else chain to return different results based on the "
            "sign of the input number."
        ),
    },
    {
        "id": 5,
        "title": "ES6 Class Practice",
        "file": ROOT / "task5.js",
        "learned": (
            "I defined an ES6 class with a constructor and a method, then created two "
            "instances with different data."
        ),
    },
]


def get_font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    preferred = [
        "/System/Library/Fonts/Menlo.ttc",
        "/System/Library/Fonts/SFNSMono.ttf",
        "/Library/Fonts/Courier New.ttf",
    ]
    for candidate in preferred:
        path = Path(candidate)
        if path.exists():
            index = 1 if bold and path.suffix.lower() == ".ttc" else 0
            try:
                return ImageFont.truetype(str(path), size=size, index=index)
            except OSError:
                continue
    return ImageFont.load_default()


def run_task(task_path: Path) -> str:
    completed = subprocess.run(
        ["node", str(task_path)],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    return completed.stdout.strip()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, width: int) -> list[str]:
    lines: list[str] = []
    for raw_line in text.splitlines():
        words = raw_line.split(" ")
        current = ""
        for word in words:
            trial = word if not current else f"{current} {word}"
            if draw.textlength(trial, font=font) <= width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = word
        lines.append(current if current else "")
    return lines


def create_console_image(task_id: int, title: str, output: str) -> Path:
    SCREENSHOT_DIR.mkdir(exist_ok=True)
    width = 1320
    height = 760
    image = Image.new("RGB", (width, height), "#111111")
    draw = ImageDraw.Draw(image)

    draw.rectangle((0, 0, width, 62), fill="#2c2c2c")
    draw.ellipse((26, 20, 42, 36), fill="#ff5f57")
    draw.ellipse((50, 20, 66, 36), fill="#febc2e")
    draw.ellipse((74, 20, 90, 36), fill="#28c840")

    title_font = get_font(26, bold=True)
    body_font = get_font(28)
    meta_font = get_font(24)

    draw.text((118, 16), f"Safari — Browser Console — Task {task_id}", fill="#f2f2f2", font=title_font)
    draw.text((40, 92), f"Homework 10 / {title}", fill="#c68d6d", font=meta_font)

    margin_x = 40
    margin_y = 150
    line_height = 42
    wrapped_lines = wrap_text(draw, output, body_font, width - (margin_x * 2))

    for index, line in enumerate(wrapped_lines):
      draw.text((margin_x, margin_y + (index * line_height)), line, fill="#f4f1ec", font=body_font)

    path = SCREENSHOT_DIR / f"task{task_id}-console.png"
    image.save(path)
    return path


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "C95F2D")
    run_properties.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)

    text_element = OxmlElement("w:t")
    text_element.text = text

    run.append(run_properties)
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)

    heading1 = document.styles["Heading 1"]
    heading1.font.name = "Aptos Display"
    heading1.font.size = Pt(18)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(0x1E, 0x1A, 0x16)

    heading2 = document.styles["Heading 2"]
    heading2.font.name = "Aptos Display"
    heading2.font.size = Pt(14)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(0x1E, 0x1A, 0x16)


def build_report() -> None:
    outputs: dict[int, str] = {}
    screenshots: dict[int, Path] = {}

    for task in TASKS:
        output = run_task(task["file"])
        outputs[task["id"]] = output
        screenshots[task["id"]] = create_console_image(task["id"], task["title"], output)

    document = Document()
    style_document(document)

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("SDT 104 – UI Design and AI-assisted Frontend Development\nHomework 10 – Introduction to JavaScript")
    title_run.bold = True
    title_run.font.size = Pt(18)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Slipchenko Illia | Spring 2026").italic = True

    document.add_paragraph("")
    links_heading = document.add_paragraph(style="Heading 1")
    links_heading.add_run("Repository & Links")

    repo_paragraph = document.add_paragraph()
    repo_paragraph.add_run("GitHub Folder: ")
    add_hyperlink(repo_paragraph, REPO_URL, REPO_URL)

    live_paragraph = document.add_paragraph()
    live_paragraph.add_run("Live Site: ")
    add_hyperlink(live_paragraph, LIVE_URL, LIVE_URL)

    note = document.add_paragraph()
    note.add_run("Note: ").bold = True
    note.add_run("Homework 10 is stored in the ")
    note.add_run("hw10").bold = True
    note.add_run(" folder of the course repository.")

    for task in TASKS:
        document.add_page_break()
        heading = document.add_paragraph(style="Heading 1")
        heading.add_run(f"Task {task['id']} – {task['title']}")

        document.add_paragraph("Code", style="Heading 2")
        code_paragraph = document.add_paragraph()
        code_run = code_paragraph.add_run(task["file"].read_text().strip())
        code_run.font.name = "Menlo"
        code_run.font.size = Pt(9.5)

        document.add_paragraph("Console Output", style="Heading 2")
        output_paragraph = document.add_paragraph()
        output_run = output_paragraph.add_run(outputs[task["id"]])
        output_run.font.name = "Menlo"
        output_run.font.size = Pt(9.5)

        document.add_paragraph("Console Screenshot", style="Heading 2")
        document.add_picture(str(screenshots[task["id"]]), width=Inches(6.8))

        document.add_paragraph("What I Learned", style="Heading 2")
        document.add_paragraph(task["learned"])

    REPORT_PATH.unlink(missing_ok=True)
    document.save(REPORT_PATH)


if __name__ == "__main__":
    build_report()
